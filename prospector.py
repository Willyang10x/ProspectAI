#!/usr/bin/env python3
import argparse
import json
import os
import re
import sys
import time
from datetime import date
from pathlib import Path
from urllib.parse import quote

import requests
from dotenv import load_dotenv
from google import genai
from google.genai import types

load_dotenv()

GOOGLE_KEY = os.getenv("GOOGLE_API_KEY", "").strip()
GEMINI_KEY = os.getenv("GEMINI_API_KEY", "").strip()

# Esconde a chave do Maps temporariamente para o SDK do Gemini não se confundir
if "GOOGLE_API_KEY" in os.environ:
    del os.environ["GOOGLE_API_KEY"]

SEU_NOME = os.getenv("SEU_NOME", "Seu Nome")
SUA_OFERTA = os.getenv(
    "SUA_OFERTA",
    "criação de sites profissionais para pequenos negócios, rápidos no celular e com botão de WhatsApp"
)
SEU_PORTFOLIO = os.getenv("SEU_PORTFOLIO", "")

DATA_FILE = Path("leads.json")
REDES_SOCIAIS = ("instagram.com", "facebook.com", "linktr.ee", "wa.me", "whatsapp.com",
                 "api.whatsapp.com", "beacons.ai", "tiktok.com", "youtube.com")

FIELDS = ",".join([
    "places.id", "places.displayName", "places.formattedAddress",
    "places.nationalPhoneNumber", "places.internationalPhoneNumber",
    "places.websiteUri", "places.rating", "places.userRatingCount",
    "places.googleMapsUri", "places.businessStatus", "nextPageToken"
])

def buscar_places(nicho: str, cidade: str, maximo: int) -> list[dict]:
    if not GOOGLE_KEY:
        sys.exit("Faltou GOOGLE_API_KEY no arquivo .env")
    url = "https://places.googleapis.com/v1/places:searchText"
    headers = {"X-Goog-Api-Key": GOOGLE_KEY, "X-Goog-FieldMask": FIELDS,
               "Content-Type": "application/json"}
    body = {"textQuery": f"{nicho} em {cidade}", "languageCode": "pt-BR",
            "regionCode": "BR", "pageSize": 20}
    
    resultados = []
    while len(resultados) < maximo:
        r = requests.post(url, headers=headers, json=body, timeout=30)
        if r.status_code != 200:
            sys.exit(f"Erro na Places API ({r.status_code}): {r.text[:300]}")
        data = r.json()
        resultados += data.get("places", [])
        token = data.get("nextPageToken")
        if not token:
            break
        body["pageToken"] = token
        time.sleep(1.5)
    return resultados[:maximo]

def pagespeed_mobile(url: str):
    try:
        r = requests.get(
            "https://www.googleapis.com/pagespeedonline/v5/runPagespeed",
            params={"url": url, "strategy": "mobile", "category": "performance",
                    **({"key": GOOGLE_KEY} if GOOGLE_KEY else {})},
            timeout=90,
        )
        score = r.json()["lighthouseResult"]["categories"]["performance"]["score"]
        return round(score * 100)
    except Exception:
        return None

def tem_whatsapp_provavel(intl_phone: str) -> str:
    digitos = re.sub(r"\D", "", intl_phone or "")
    if digitos.startswith("55") and len(digitos) == 13 and digitos[4] == "9":
        return digitos
    return ""

def qualificar(p: dict, usar_pagespeed: bool) -> dict:
    nome = p.get("displayName", {}).get("text", "")
    site = p.get("websiteUri", "") or ""
    nota = p.get("rating")
    avals = p.get("userRatingCount", 0) or 0
    zap = tem_whatsapp_provavel(p.get("internationalPhoneNumber", ""))
    
    score, motivos = 0, []
    site_baixo = site.lower()
    so_rede_social = any(d in site_baixo for d in REDES_SOCIAIS)
    ps = None
    
    if not site:
        score += 50
        motivos.append("sem site próprio")
    elif so_rede_social:
        score += 45
        motivos.append("usa só rede social/link como site")
    else:
        if site.startswith("http://"):
            score += 15
            motivos.append("site sem HTTPS")
        if usar_pagespeed:
            ps = pagespeed_mobile(site)
            if ps is not None and ps < 50:
                score += 20
                motivos.append(f"site lento no celular (nota {ps}/100)")
            elif ps is not None and ps < 75:
                score += 8
                motivos.append(f"site mediano no celular (nota {ps}/100)")
                
    if avals >= 30:
        score += 15
        motivos.append(f"{avals} avaliações (negócio ativo)")
    if nota and nota >= 4.3:
        score += 10
        motivos.append(f"nota {nota} no Google")
    if zap:
        score += 5
        
    if p.get("businessStatus") not in (None, "OPERATIONAL"):
        score = 0
        motivos.append("não está operacional")
        
    return {
        "id": p["id"], "nome": nome, "endereco": p.get("formattedAddress", ""),
        "telefone": p.get("nationalPhoneNumber", ""), "whatsapp": zap,
        "site": site, "nota": nota, "avaliacoes": avals,
        "maps": p.get("googleMapsUri", ""), "pagespeed": ps,
        "score": min(score, 100), "diagnostico": "; ".join(motivos) or "site ok",
        "mensagens": {},
    }

SYSTEM_PROMPT = f"""Você é {SEU_NOME}, desenvolvedor(a) que oferece {SUA_OFERTA}. Escreva abordagens comerciais frias em português do Brasil para um negócio local. 
Regras:
- Use SOMENTE os dados fornecidos; nunca invente fatos sobre o negócio.
- Cite 1 detalhe real (nota, avaliações, ausência de site, etc.) de forma elogiosa e natural.
- Tom humano, respeitoso e direto, sem parecer spam. Sem exageros, sem emoji em excesso (no máximo 1).
- Adapte ao nicho (advocacia pede tom mais sóbrio e discreto; clínicas, acolhedor).
- WhatsApp: até 450 caracteres, termina com pergunta simples. Inclua uma saída gentil ("se não fizer sentido, é só avisar que não incomodo mais").
- Instagram DM: até 350 caracteres.
- E-mail: assunto curto + corpo de 90-130 palavras.
- followup: mensagem de WhatsApp de 3 a 5 dias depois, curta e leve.
{('- Se couber, mencione o portfólio: ' + SEU_PORTFOLIO) if SEU_PORTFOLIO else ''}

Responda APENAS com JSON válido, com as seguintes chaves exatas: "whatsapp", "instagram", "email_assunto", "email_corpo", "followup" """

def mensagem_fallback(l: dict) -> dict:
    nome = l["nome"]
    gancho = ("Vi que o(a) {n} tem {a} avaliações no Google" if l["avaliacoes"] >= 10
              else "Encontrei o(a) {n} pesquisando no Google").format(n=nome, a=l["avaliacoes"])
    dor = ("mas não achei um site próprio" if not l["site"] or
           any(d in l["site"].lower() for d in REDES_SOCIAIS)
           else "e dei uma olhada no site de vocês, vejo espaço para melhorar")
           
    zap = (f"Olá! Aqui é {SEU_NOME}. {gancho}, {dor}. Trabalho com {SUA_OFERTA}. "
           "Posso te mostrar uma ideia rápida, sem compromisso? "
           "Se não fizer sentido, é só avisar que não incomodo mais.")
           
    return {
        "whatsapp": zap,
        "instagram": zap[:350],
        "email_assunto": f"Uma ideia para a presença online do(a) {nome}",
        "email_corpo": f"Olá, tudo bem?\n\n{gancho}, {dor}.\n\nTrabalho com {SUA_OFERTA} e "
                       f"gostaria de mostrar uma proposta enxuta.\n\nPosso enviar? "
                       f"Se não for o momento, é só me avisar.\n\nAbraço,\n{SEU_NOME}",
        "followup": f"Oi! Só retomando minha mensagem sobre o site do(a) {nome}. "
                    "Faz sentido eu te mandar uma ideia? Se não, sem problemas!",
    }

def gerar_mensagens(l: dict, nicho: str, cliente) -> dict:
    if cliente is None:
        return mensagem_fallback(l)
        
    dados = {k: l[k] for k in ("nome", "endereco", "site", "nota", "avaliacoes", "diagnostico")}
    dados["nicho"] = nicho
    prompt = json.dumps(dados, ensure_ascii=False)
    
    tentativas = 3
    for t in range(tentativas):
        try:
            resp = cliente.models.generate_content(
                model="gemini-2.5-flash",
                contents=prompt,
                config=types.GenerateContentConfig(
                    system_instruction=SYSTEM_PROMPT,
                    response_mime_type="application/json"
                )
            )
            msgs = json.loads(resp.text)
            
            if not all(k in msgs for k in ("whatsapp", "instagram", "email_assunto", "email_corpo", "followup")):
                raise ValueError("JSON incompleto")
                
            return msgs
            
        except Exception as e:
            # Se for erro de tráfego, espera e tenta de novo
            if "503" in str(e) and t < tentativas - 1:
                print(f"   - Servidor do Gemini ocupado. Tentando novamente em 3 segundos... ({t+1}/{tentativas})")
                time.sleep(3)
            else:
                print(f"   ! IA falhou para {l['nome']} ({e}); usando modelo padrão.")
                return mensagem_fallback(l)

def wa_link(l: dict, texto: str) -> str:
    return f"https://wa.me/{l['whatsapp']}?text={quote(texto)}" if l["whatsapp"] else ""

def exportar_xlsx(leads: list[dict], caminho: str = "leads.xlsx"):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    
    cab = ["Score", "Nome", "Telefone", "WhatsApp provável", "Site", "Nota", "Avaliações",
           "Diagnóstico", "Mensagem WhatsApp", "Link WhatsApp", "Mensagem Instagram",
           "Assunto e-mail", "Corpo e-mail", "Follow-up", "Google Maps", "Status"]
    ws.append(cab)
    
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        
    for l in leads:
        m = l.get("mensagens") or {}
        ws.append([l["score"], l["nome"], l["telefone"], "sim" if l["whatsapp"] else "não",
                   l["site"] or "(sem site)", l["nota"], l["avaliacoes"], l["diagnostico"],
                   m.get("whatsapp", ""), wa_link(l, m.get("whatsapp", "")),
                   m.get("instagram", ""), m.get("email_assunto", ""), m.get("email_corpo", ""),
                   m.get("followup", ""), l["maps"], l.get("status", "novo")])
                   
    larguras = [7, 34, 16, 12, 30, 7, 11, 40, 60, 30, 50, 30, 60, 45, 30, 12]
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[chr(64 + i)].width = w
        
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(caminho)

HTML = """<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"> <meta name="viewport" content="width=device-width,initial-scale=1"><title>Painel de Prospecção</title> <style> :root{--bg:#f5f6f8;--card:#fff;--tx:#1b1f23;--mut:#6b7280;--bd:#e3e6ea;--ac:#1f3864} @media(prefers-color-scheme:dark){:root{--bg:#14171a;--card:#1d2125;--tx:#e8eaed;--mut:#9aa0a6;--bd:#2d3237;--ac:#7aa2f7}} *{box-sizing:border-box}body{margin:0;font:15px/1.45 system-ui,Arial,sans-serif;background:var(--bg);color:var(--tx);padding:16px;max-width:980px;margin:auto} h1{font-size:20px}.bar{display:flex;gap:8px;flex-wrap:wrap;margin:12px 0} input,select{padding:8px 10px;border:1px solid var(--bd);border-radius:8px;background:var(--card);color:var(--tx)} .card{background:var(--card);border:1px solid var(--bd);border-radius:12px;padding:14px;margin-bottom:12px} .top{display:flex;justify-content:space-between;gap:10px;align-items:flex-start} .nome{font-weight:600;font-size:16px}.mut{color:var(--mut);font-size:13px} .score{background:var(--ac);color:#fff;border-radius:8px;padding:2px 9px;font-weight:600;white-space:nowrap} .msg{white-space:pre-wrap;background:var(--bg);border-radius:8px;padding:10px;margin:8px 0;font-size:14px} .btns{display:flex;gap:6px;flex-wrap:wrap}button,a.b{padding:7px 11px;border-radius:8px;border:1px solid var(--bd);background:var(--card);color:var(--tx);cursor:pointer;text-decoration:none;font-size:13px} a.wa{background:#25d366;color:#052e16;border-color:#25d366;font-weight:600} details summary{cursor:pointer;color:var(--ac);margin-top:6px} </style></head><body> <h1>Painel de Prospecção <span class="mut" id="cont"></span></h1> <div class="bar"><input id="q" placeholder="Buscar nome..." oninput="draw()"> <select id="st" onchange="draw()"><option value="">Todos os status</option><option>novo</option><option>enviado</option><option>respondeu</option><option>reunião</option><option>fechado</option><option>descartado</option></select> <select id="zap" onchange="draw()"><option value="">Com ou sem WhatsApp</option><option value="1">Só com WhatsApp provável</option></select></div> <div id="list"></div> <script> const LEADS=__DATA__; let S={};try{S=JSON.parse(localStorage.getItem('status')||'{}')}catch(e){} const esc=s=>String(s??'').replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c])); function setSt(id,v){S[id]=v;try{localStorage.setItem('status',JSON.stringify(S))}catch(e){}draw()} function cp(t,b){navigator.clipboard.writeText(t).then(()=>{const o=b.textContent;b.textContent='Copiado!';setTimeout(()=>b.textContent=o,1200)})} function draw(){  const q=document.getElementById('q').value.toLowerCase(),st=document.getElementById('st').value,z=document.getElementById('zap').value;  const L=LEADS.filter(l=>l.nome.toLowerCase().includes(q)&&(!st||(S[l.id]||'novo')==st)&&(!z||l.whatsapp));  document.getElementById('cont').textContent='('+L.length+' leads)';  document.getElementById('list').innerHTML=L.map((l,i)=>{const m=l.mensagens||{},s=S[l.id]||'novo';  return `<div class="card"><div class="top"><div><div class="nome">${esc(l.nome)}</div>  <div class="mut">${esc(l.endereco)}<br>${esc(l.telefone)} • ${l.nota?('⭐ '+l.nota+' ('+l.avaliacoes+')'):'sem nota'} • ${l.site?`<a href="${esc(l.site)}" target="_blank" rel="noopener">${esc(l.site)}</a>`:'sem site'}</div>  <div class="mut"><b>Diagnóstico:</b> ${esc(l.diagnostico)}</div></div><span class="score">${l.score}</span></div>  <div class="msg">${esc(m.whatsapp||'(sem mensagem gerada, rode com --top maior)')}</div>  <div class="btns">  ${l.wa_link?`<a class="b wa" href="${esc(l.wa_link)}" target="_blank" rel="noopener">Abrir no WhatsApp</a>`:''}  <button onclick="cp(LEADS_BY[${i}].mensagens.whatsapp,this)">Copiar WhatsApp</button>  <button onclick="cp(LEADS_BY[${i}].mensagens.instagram,this)">Copiar Instagram</button>  <a class="b" href="${esc(l.maps)}" target="_blank" rel="noopener">Maps</a>  <select onchange="setSt('${l.id}',this.value)">${['novo','enviado','respondeu','reunião','fechado','descartado'].map(o=>`<option ${o==s?'selected':''}>${o}</option>`).join('')}</select></div>  <details><summary>E-mail e follow-up</summary><div class="msg"><b>${esc(m.email_assunto)}</b>\\n\\n${esc(m.email_corpo)}</div>  <div class="msg"><b>Follow-up (3-5 dias):</b>\\n${esc(m.followup)}</div>  <div class="btns"><button onclick="cp(LEADS_BY[${i}].mensagens.email_corpo,this)">Copiar e-mail</button><button onclick="cp(LEADS_BY[${i}].mensagens.followup,this)">Copiar follow-up</button></div></details></div>`}).join('');  window.LEADS_BY=L} draw(); </script></body></html>"""

def exportar_html(leads: list[dict], caminho: str = "painel.html"):
    for l in leads:
        l["wa_link"] = wa_link(l, (l.get("mensagens") or {}).get("whatsapp", ""))
    dados = json.dumps(leads, ensure_ascii=False).replace("</", "<\\/")
    Path(caminho).write_text(HTML.replace("__DATA__", dados), encoding="utf-8")

def main():
    ap = argparse.ArgumentParser(description="Agente de prospecção via Google Maps")
    ap.add_argument("nicho", help='ex: "advogado", "clínica odontológica"')
    ap.add_argument("cidade", help='ex: "Campina Grande PB"')
    ap.add_argument("--max", type=int, default=40, help="máx. de lugares a buscar")
    ap.add_argument("--top", type=int, default=20, help="gerar mensagens só para os N melhores leads novos")
    ap.add_argument("--pagespeed", action="store_true", help="testar velocidade dos sites (mais lento)")
    ap.add_argument("--sem-ia", action="store_true", help="usar só mensagens-modelo (sem custo de API)")
    args = ap.parse_args()
    
    base = {l["id"]: l for l in json.loads(DATA_FILE.read_text("utf-8"))} if DATA_FILE.exists() else {}
    
    print(f"Buscando '{args.nicho}' em {args.cidade}...")
    lugares = buscar_places(args.nicho, args.cidade, args.max)
    print(f"{len(lugares)} lugares encontrados. Qualificando...")
    
    novos = []
    for p in lugares:
        if p["id"] in base:
            continue
        l = qualificar(p, args.pagespeed)
        l.update(nicho=args.nicho, cidade=args.cidade, data=str(date.today()))
        novos.append(l)
        
    novos.sort(key=lambda x: x["score"], reverse=True)
    
    cliente = None
    if GEMINI_KEY and not args.sem_ia:
        cliente = genai.Client(api_key=GEMINI_KEY)
    else:
        print("Sem chave do Gemini (ou --sem-ia): usando mensagens-modelo.")
        
    for i, l in enumerate(novos[: args.top], 1):
        print(f"[{i}/{min(args.top, len(novos))}] Mensagens para {l['nome']} (score {l['score']})")
        l["mensagens"] = gerar_mensagens(l, args.nicho, cliente)
        
    for l in novos[args.top:]:
        l["mensagens"] = {}
        
    for l in novos:
        base[l["id"]] = l
        
    todos = sorted(base.values(), key=lambda x: x["score"], reverse=True)
    DATA_FILE.write_text(json.dumps(todos, ensure_ascii=False, indent=1), "utf-8")
    exportar_xlsx(todos)
    exportar_html(todos)
    
    print(f"\nPronto! {len(novos)} leads novos ({len(todos)} no total).")
    print("Abra painel.html no navegador ou leads.xlsx no Excel.")

if __name__ == "__main__":
    main()